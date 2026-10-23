"""稿间: a local, single-user writing workflow prototype."""
import asyncio
from contextlib import asynccontextmanager
from datetime import datetime
from io import BytesIO
import os
from pathlib import Path
import shutil
import uuid
from urllib.parse import urlsplit
from typing import Literal
from zoneinfo import ZoneInfo

from docx import Document
from fastapi import FastAPI, File, Form, HTTPException, UploadFile
from fastapi.responses import Response
from fastapi.staticfiles import StaticFiles
from PIL import Image, UnidentifiedImageError
from pydantic import BaseModel, Field

import engine
import providers
from store import Store

ROOT = Path(__file__).resolve().parent
DATA = Path(os.environ.get('WRITING_STUDIO_DATA', str(ROOT / 'data'))).resolve()
TZ = ZoneInfo('Asia/Shanghai')
ACTIVE = {'queued', 'extracting', 'writing', 'reviewing', 'revising'}
DEFAULT_SETTINGS = {'base_url': 'https://api.openai.com/v1', 'api_key': '',
                    'writer_model': '', 'reviewer_model': '', 'vision_model': '', 'daily_limit': 100000,
                    'reviewer_base_url': '', 'reviewer_api_key': ''}
DEFAULT_PROFILE = {'pen_name': '', 'style': '自然、具体、克制，减少套话与空泛议论。', 'materials': ''}


def now():
    return datetime.now(TZ).isoformat(timespec='seconds')


def public_task(task):
    return {k: v for k, v in task.items() if not k.startswith('_')}


def new_task(title, source_type, source_text='', source_url='', mode='demo', max_rounds=2,
             author_notes='', fiction=True, is_sample=False):
    return {'id': uuid.uuid4().hex[:12], 'title': title or '新的征稿任务', 'status': 'queued',
            'stage': 'queued', 'message': '等待开始处理', 'mode': mode,
            'source_type': source_type, 'source_text': source_text, 'source_url': source_url,
            'requirements': {}, 'draft': '', 'word_count': 0, 'round': 0,
            'max_rounds': max_rounds, 'reviews': [], 'versions': [], 'checks': [],
            'events': [], 'blockers': [], 'created_at': now(), 'updated_at': now(),
            'is_sample': is_sample, 'fiction': fiction, 'author_notes': author_notes,
            '_action': 'create', '_images': [], '_reviewed_version': -1}


@asynccontextmanager
async def lifespan(app):
    app.state.store = Store(DATA)
    app.state.reserved = 0
    if not app.state.store.get_config('seeded', False):
        if os.environ.get('WRITING_STUDIO_SEED', '1') == '1':
            task = new_task('示例征稿 · 日常里的微光', 'text', engine.demo_source(), is_sample=True)
            app.state.store.save(task)
        app.state.store.set_config('seeded', True)
    # Completed steps are retained; the worker resumes the first incomplete step.
    for task in app.state.store.all():
        if task['status'] in ACTIVE:
            task.update(status='queued', message='已恢复任务，继续未完成的步骤')
            app.state.store.save(task)
    worker = asyncio.create_task(work_loop(app))
    yield
    worker.cancel()
    try:
        await worker
    except asyncio.CancelledError:
        pass
    app.state.store.close()


app = FastAPI(title='稿间 · 写作工作台', lifespan=lifespan, docs_url='/api/docs', redoc_url=None)


@app.middleware('http')
async def local_protection(request, call_next):
    # The prototype is single-user. Reject cross-site browser writes, even locally.
    if request.method not in {'GET', 'HEAD', 'OPTIONS'}:
        if request.headers.get('sec-fetch-site') == 'cross-site':
            return Response('Cross-site writes are disabled', status_code=403)
        origin = request.headers.get('origin')
        if origin and urlsplit(origin).netloc != request.headers.get('host'):
            return Response('Cross-origin writes are disabled', status_code=403)
    response = await call_next(request)
    response.headers['X-Content-Type-Options'] = 'nosniff'
    response.headers['Referrer-Policy'] = 'no-referrer'
    response.headers['Content-Security-Policy'] = "default-src 'self'; script-src 'self'; style-src 'self' 'unsafe-inline'; img-src 'self' data: blob:; connect-src 'self'; frame-ancestors 'none'"
    if request.url.path.startswith('/api/'):
        response.headers['Cache-Control'] = 'no-store'
    return response


def store():
    return app.state.store


def get_task(task_id):
    task = store().get(task_id)
    if not task:
        raise HTTPException(404, '任务不存在')
    return task


def settings():
    saved = {**DEFAULT_SETTINGS, **store().get_config('settings', {})}
    if not saved['api_key']:
        saved['api_key'] = os.environ.get('WRITING_API_KEY', '')
    for name in ('writer_model', 'reviewer_model', 'vision_model'):
        saved[name] = str(saved.get(name) or '').strip()
    if not saved['reviewer_base_url'].strip() and not saved['reviewer_model']:
        saved['reviewer_model'] = saved['writer_model']
    return saved


def profile():
    return {**DEFAULT_PROFILE, **store().get_config('profile', {})}


def ledger():
    current = store().get_config('usage', {})
    if current.get('day') != now()[:10]:
        current = {'day': now()[:10], 'used_tokens': 0, 'estimated': False, 'calls': 0}
    return current


def configuration_gaps(data):
    missing = []
    if not data['api_key']:
        missing.append('API密钥')
    if not data['writer_model']:
        missing.append('写作模型')
    if data['reviewer_base_url'].strip() and not data['reviewer_model']:
        missing.append('评审模型（独立评审服务）')
    return missing


def configuration_message():
    missing = configuration_gaps(settings())
    return '模型设置尚未完成，请填写并保存：' + '、'.join(missing) + '。尚未向服务商验证密钥。'


def public_settings():
    data = settings()
    missing = configuration_gaps(data)
    return {k: v for k, v in data.items() if k not in {'api_key', 'reviewer_api_key'}} | {
        'has_key': bool(data['api_key']),
        'has_reviewer_key': bool(data['reviewer_api_key']),
        'configured': not missing, 'missing_fields': missing,
        'used_tokens': ledger()['used_tokens'], 'usage_estimated': ledger()['estimated']}


def call_settings():
    data = settings()

    def reserve(amount):
        if ledger()['used_tokens'] + app.state.reserved + amount > data['daily_limit']:
            raise engine.EngineError('今日 token 上限不足以开始下一次调用，请提高设置中的上限或次日继续。')
        app.state.reserved += amount

    def release(amount):
        app.state.reserved = max(0, app.state.reserved - amount)

    def charge(usage):
        current = ledger()
        current['used_tokens'] += max(0, int(usage.get('total_tokens', 0)))
        current['estimated'] = current['estimated'] or bool(usage.get('estimated'))
        current['calls'] += 1
        store().set_config('usage', current)

    data.update(_reserve_callback=reserve, _release_callback=release, _usage_callback=charge)
    return data


def event(task, stage, message):
    task.update(status=stage, stage=stage, message=message, updated_at=now())
    task['events'].append({'stage': stage, 'message': message, 'created_at': now()})
    task['events'] = task['events'][-60:]
    store().save(task)


def save_version(task, title):
    task['versions'].append({'round': len(task['versions']), 'title': title,
                             'content': task['draft'], 'created_at': now()})
    task['word_count'] = engine.count_words(task['draft'])
    store().save(task)


def preflight(task):
    req = task['requirements']
    blockers = []
    original = engine.extract_local_requirements(task['source_text'])
    if original.get('multiple_notices'):
        blockers.append('材料中包含多条征稿，请在补充正文中保留目标征稿的完整段落，避免混合不同要求。')
    if req.get('ai_policy') == 'forbidden' or original.get('ai_policy') == 'forbidden':
        blockers.append('征稿明确禁止 AI 创作，已停止自动写作。可以保留信息，另行独立创作。')
    needs_facts = req.get('requires_true_experience') or original.get('requires_true_experience')
    if needs_facts and (task['fiction'] or not (task['author_notes'].strip() or profile()['materials'].strip())):
        blockers.append('该征稿需要真实经历：请选择真实素材创作并补充亲身素材，系统不会代编经历。')
    date = req.get('deadline')
    if date:
        try:
            if datetime.fromisoformat(str(date)[:10]).date() < datetime.now(TZ).date():
                blockers.append('征稿截止日期已过，请核对原文或选择新的征稿。')
        except ValueError:
            pass  # Unspecified/ambiguous dates are shown as warnings, never guessed.
    return blockers


def draft_checks(task):
    materials = '\n'.join(filter(None, [task['author_notes'], profile()['materials']]))
    checks = engine.check_draft(task['draft'], task['requirements'], fiction=task['fiction'], author_notes=materials)
    for blocker in preflight(task):
        checks.append({'label': '写作前提', 'status': 'fail', 'detail': blocker})
    if task['mode'] == 'demo':
        checks.append({'label': '演示内容', 'status': 'warn', 'detail': '本稿为本地规则生成的演示样稿，未调用真实模型，请勿直接投稿。'})
    return checks


def needs_revision(task):
    if any(c['status'] == 'fail' for c in task['checks']):
        return True
    for review in task['reviews']:
        for issue in review.get('issues', []):
            if issue.get('severity') in {'critical', 'high', 'medium', '严重', '中等'}:
                return True
    return False


async def review_current(task, config):
    version = len(task['versions']) - 1
    if task.get('_reviewed_version') == version:
        return
    event(task, 'reviewing', '两路独立评审正在核对要求、结构与表达')
    task['reviews'] = await engine.review_draft(task['draft'], task['requirements'], settings=config, mode=task['mode'])
    task['_reviewed_version'] = version
    task['checks'] = draft_checks(task)
    store().save(task)


async def process_task(task):
    config = call_settings()
    if task['mode'] == 'live' and not public_settings()['configured']:
        raise engine.EngineError(configuration_message())
    if not task['source_text'].strip():
        event(task, 'extracting', '正在读取征稿来源')
        try:
            if task['source_type'] == 'url':
                task['source_text'] = await engine.fetch_source(task['source_url'])
            elif task['source_type'] == 'images':
                if task['mode'] != 'live' or not config.get('vision_model') or not config.get('api_key'):
                    raise engine.EngineError('截图识别需要真实视觉模型。请配置后重试，或在此任务中补充征稿正文。')
                images = [(Path(p).read_bytes(), mime) for p, mime in task['_images']]
                task['source_text'] = await engine.ocr_images(images, settings=config)
            store().save(task)
        except engine.EngineError as exc:
            task['blockers'] = [str(exc)]
            event(task, 'needs_input', str(exc))
            return
    if not task['source_text'].strip():
        task['blockers'] = ['未取得可识别的正文，请补充征稿内容。']
        event(task, 'needs_input', task['blockers'][0])
        return
    if not task['requirements']:
        event(task, 'extracting', '正在提取主题、字数、截止日期与 AI 使用约定')
        task['requirements'] = await engine.extract_requirements(task['source_text'], settings=config, mode=task['mode'], fiction=task['fiction'], author_notes=task['author_notes'])
        if task['title'] == '新的征稿任务' and task['requirements'].get('title'):
            task['title'] = task['requirements']['title'][:100]
        store().save(task)
    task['blockers'] = preflight(task)
    if task['blockers']:
        event(task, 'needs_input', task['blockers'][0])
        return
    if task['_action'] == 'refine':
        # The instruction is consumed only when a new version is durably saved.
        event(task, 'revising', '正在按你的定点意见修改')
        task['draft'] = await engine.revise_draft(task['draft'], task['requirements'], task['reviews'], profile=profile(), author_notes=task['author_notes'], instruction=task.get('_instruction', ''), fiction=task['fiction'], settings=config, mode=task['mode'])
        task['_action'] = 'finish_refine'
        save_version(task, '按作者意见修改')
    if not task['draft']:
        event(task, 'writing', '正在结合征稿要求与作者素材形成初稿')
        task['draft'] = await engine.write_draft(task['requirements'], source_text=task['source_text'], profile=profile(), author_notes=task['author_notes'], fiction=task['fiction'], settings=config, mode=task['mode'])
        save_version(task, '初稿')
    await review_current(task, config)
    while task['_action'] == 'create' and needs_revision(task) and task['round'] < task['max_rounds']:
        event(task, 'revising', f"自动修改第 {task['round'] + 1} 轮：综合评审意见")
        previous = task['draft']
        task['draft'] = await engine.revise_draft(previous, task['requirements'], task['reviews'], profile=profile(), author_notes=task['author_notes'], instruction='', fiction=task['fiction'], settings=config, mode=task['mode'])
        task['round'] += 1
        save_version(task, f"自动修改 · 第 {task['round']} 轮")
        await review_current(task, config)
        if task['draft'] == previous:
            break
    task['_action'] = 'complete'
    task['checks'] = draft_checks(task)
    task['blockers'] = [check['detail'] for check in task['checks'] if check['status'] == 'fail']
    message = '审阅包已生成，请核对征稿原文、终稿和待确认事项'
    if task['blockers']:
        message = '自动处理已结束，仍有硬性要求待处理'
    event(task, 'needs_review', message)


async def work_loop(application):
    while True:
        candidates = [task for task in application.state.store.all() if task['status'] == 'queued']
        if not candidates:
            await asyncio.sleep(0.25)
            continue
        task = candidates[-1]
        try:
            await process_task(task)
        except asyncio.CancelledError:
            raise
        except Exception as exc:
            if isinstance(exc, engine.EngineError):
                message = str(exc)[:600]
            else:
                message = f'当前步骤处理失败（{type(exc).__name__}），已保留已完成内容，可稍后重试。'
            task['blockers'] = [message]
            event(task, 'failed', message)


@app.get('/api/health')
def health():
    return {'status': 'ok', 'service': '稿间'}


@app.get('/api/bootstrap')
def bootstrap():
    tasks = [public_task(task) for task in store().all()]
    return {'tasks': tasks, 'profile': profile(), 'settings': public_settings(), 'providers': providers.PROVIDERS, 'mode': 'demo',
            'stats': {'total': len(tasks), 'active': sum(t['status'] in ACTIVE for t in tasks),
                      'review': sum(t['status'] == 'needs_review' for t in tasks),
                      'approved': sum(t['status'] == 'approved' for t in tasks)}}


@app.get('/api/tasks')
def tasks():
    return [public_task(task) for task in store().all()]


@app.get('/api/providers')
def provider_presets():
    return providers.PROVIDERS


class ModelRole(BaseModel):
    role: Literal['writer', 'reviewer', 'vision'] = 'writer'


@app.post('/api/settings/models')
async def available_models(body: ModelRole):
    try:
        return await providers.list_models(settings(), body.role)
    except engine.EngineError as exc:
        raise HTTPException(400, str(exc))


@app.post('/api/settings/test')
async def check_connection(body: ModelRole):
    try:
        return await providers.test_connection(call_settings(), body.role)
    except engine.EngineError as exc:
        raise HTTPException(400, str(exc))


@app.get('/api/tasks/{task_id}')
def task_detail(task_id: str):
    return public_task(get_task(task_id))


@app.post('/api/tasks', status_code=202)
async def create_task(title: str = Form(''), source_type: str = Form('text'), content: str = Form(''),
                      url: str = Form(''), mode: str = Form('demo'), max_rounds: int = Form(2),
                      author_notes: str = Form(''), fiction: bool = Form(True),
                      images: list[UploadFile] = File(default=[])):
    if mode not in {'demo', 'live'} or source_type not in {'text', 'url', 'images'}:
        raise HTTPException(422, '不支持的来源或运行模式')
    if not 1 <= max_rounds <= 3:
        raise HTTPException(422, '自动修改上限为 1 至 3 轮')
    if len(title) > 100 or len(content) > 30000 or len(author_notes) > 12000 or len(url) > 2048:
        raise HTTPException(422, '输入过长，请缩短正文或素材后重试')
    if source_type == 'text' and len(content.strip()) < 10:
        raise HTTPException(422, '请粘贴至少 10 个字符的征稿正文')
    if source_type == 'url' and not url.strip():
        raise HTTPException(422, '请填写公众号文章链接')
    if source_type == 'images' and not 1 <= len(images) <= 8:
        raise HTTPException(422, '请上传 1 至 8 张截图')
    if mode == 'live' and not public_settings()['configured']:
        raise HTTPException(400, configuration_message())
    task = new_task(title.strip(), source_type, content.strip() if source_type == 'text' else '',
                    url.strip(), mode, max_rounds, author_notes.strip(), fiction)
    image_directory = DATA / 'uploads' / task['id']
    try:
        if source_type == 'images':
            image_directory.mkdir(parents=True, mode=0o700)
            for index, upload in enumerate(images):
                body = await upload.read(5 * 1024 * 1024 + 1)
                if len(body) > 5 * 1024 * 1024:
                    raise HTTPException(413, '单张截图不能超过 5 MB')
                try:
                    with Image.open(BytesIO(body)) as image:
                        if image.format not in {'JPEG', 'PNG', 'WEBP'} or image.width * image.height > 20000000:
                            raise ValueError('Unsupported image')
                        mime = Image.MIME[image.format]
                        image.verify()
                except (UnidentifiedImageError, ValueError, OSError, Image.DecompressionBombError):
                    raise HTTPException(422, '截图须为有效的 JPG、PNG 或 WebP，且不超过 2000 万像素')
                path = image_directory / str(index)
                path.write_bytes(body)
                task['_images'].append((str(path), mime))
        store().save(task)
    except Exception:
        if image_directory.exists():
            shutil.rmtree(image_directory)
        raise
    return public_task(task)


class Edit(BaseModel):
    draft: str = Field(min_length=1, max_length=30000)


def editable(task):
    if task['status'] in ACTIVE:
        raise HTTPException(409, '任务正在处理中，请稍后操作')


@app.post('/api/tasks/{task_id}/edit')
def edit_task(task_id: str, body: Edit):
    task = get_task(task_id)
    editable(task)
    if not task['requirements'] or not task['draft']:
        raise HTTPException(409, '请先完成初稿生成')
    task['draft'] = body.draft
    save_version(task, '作者手动编辑')
    task['checks'] = draft_checks(task)
    task['blockers'] = [c['detail'] for c in task['checks'] if c['status'] == 'fail']
    task['reviews'] = []
    task['_reviewed_version'] = -1
    event(task, 'needs_review', '手动编辑已保存，原 AI 评审已失效，请核对修改后的终稿')
    return public_task(task)


class Refine(BaseModel):
    instruction: str = Field(min_length=2, max_length=3000)


@app.post('/api/tasks/{task_id}/refine', status_code=202)
def refine_task(task_id: str, body: Refine):
    task = get_task(task_id)
    editable(task)
    if not task['draft']:
        raise HTTPException(409, '当前还没有可修改的初稿')
    task.update(_action='refine', _instruction=body.instruction)
    event(task, 'queued', '已收到定点修改意见，等待修改与重新评审')
    return public_task(task)


class Approval(BaseModel):
    acknowledged: bool = False


@app.post('/api/tasks/{task_id}/approve')
def approve_task(task_id: str, body: Approval):
    task = get_task(task_id)
    if task['status'] != 'needs_review':
        raise HTTPException(409, '只有待审阅的稿件可以确认')
    if not body.acknowledged:
        raise HTTPException(422, '请先确认已核对原文、真实事实和 AI 使用要求')
    task['checks'] = draft_checks(task)
    if any(check['status'] == 'fail' for check in task['checks']):
        raise HTTPException(409, '仍有硬性要求未满足，请先修改稿件或核实征稿内容')
    event(task, 'approved', '作者已完成审阅；尚未发送或投稿')
    task['approved_at'] = now()
    store().save(task)
    return public_task(task)


class Retry(BaseModel):
    content: str | None = Field(default=None, max_length=30000)
    author_notes: str | None = Field(default=None, max_length=12000)
    fiction: bool | None = None


@app.post('/api/tasks/{task_id}/retry', status_code=202)
def retry_task(task_id: str, body: Retry):
    task = get_task(task_id)
    editable(task)
    if body.content is not None:
        if len(body.content.strip()) < 10:
            raise HTTPException(422, '请补充至少 10 个字符的征稿正文')
        task.update(source_text=body.content.strip(), source_type='text', requirements={}, draft='',
                    reviews=[], checks=[], versions=[], round=0, _reviewed_version=-1, _action='create')
    if body.author_notes is not None:
        task['author_notes'] = body.author_notes
    if body.fiction is not None:
        task['fiction'] = body.fiction
    if task['_action'] == 'complete':
        task['_action'] = 'create'
    task['blockers'] = []
    event(task, 'queued', '已补充信息，继续处理未完成的步骤')
    return public_task(task)


@app.delete('/api/tasks/{task_id}')
def delete_task(task_id: str):
    task = get_task(task_id)
    editable(task)
    store().delete(task_id)
    directory = DATA / 'uploads' / task_id
    if directory.exists():
        shutil.rmtree(directory)
    return {'deleted': True}


@app.get('/api/tasks/{task_id}/export')
def export(task_id: str, format: str = 'txt'):
    task = get_task(task_id)
    if not task['draft']:
        raise HTTPException(409, '尚无稿件可导出')
    title = task['title']
    warning = '【演示样稿，未调用真实模型，请勿直接投稿】\n\n' if task['mode'] == 'demo' else ''
    if format == 'email':
        req = task['requirements']
        text = (f"{warning}收件人：{req.get('submission') or '请核实征稿原文中的地址'}\n"
                f"主题：投稿｜{title}｜{profile()['pen_name'] or '请填写笔名'}\n\n"
                '编辑老师您好：\n\n随信附上稿件，请您审阅。\n\n'
                f"{task['draft']}\n\n作者：{profile()['pen_name'] or '请填写'}\n"
                '作者简介与联系方式：请在发送前补充并核实。\n\n【本文件仅为邮件草稿，未执行发送】')
        body, media, extension = text.encode('utf-8-sig'), 'text/plain; charset=utf-8', 'txt'
    elif format == 'txt':
        body, media, extension = (warning + task['draft']).encode('utf-8-sig'), 'text/plain; charset=utf-8', 'txt'
    elif format == 'docx':
        document = Document()
        document.add_heading(title, 0)
        if warning:
            document.add_paragraph(warning.strip())
        for paragraph in task['draft'].split('\n'):
            if paragraph.strip():
                document.add_paragraph(paragraph.strip())
        buffer = BytesIO()
        document.save(buffer)
        body, media, extension = buffer.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document', 'docx'
    else:
        raise HTTPException(422, '支持导出 txt、docx 或 email')
    return Response(body, media_type=media, headers={'Content-Disposition': f'attachment; filename="gaojian-{task_id}.{extension}"'})


class ProfileInput(BaseModel):
    pen_name: str = Field(default='', max_length=50)
    style: str = Field(default='', max_length=2000)
    materials: str = Field(default='', max_length=12000)


@app.put('/api/profile')
def save_profile(body: ProfileInput):
    store().set_config('profile', body.model_dump())
    return profile()


class SettingsInput(BaseModel):
    base_url: str = Field(max_length=300)
    api_key: str | None = Field(default=None, max_length=500)
    reviewer_base_url: str = Field(default='', max_length=300)
    reviewer_api_key: str | None = Field(default=None, max_length=500)
    writer_model: str = Field(default='', max_length=150)
    reviewer_model: str = Field(default='', max_length=150)
    vision_model: str = Field(default='', max_length=150)
    daily_limit: int = Field(default=100000, ge=1000, le=10000000)


@app.put('/api/settings')
def save_settings(body: SettingsInput):
    for key in (body.api_key, body.reviewer_api_key):
        if key and any(marker * 3 in key for marker in ('*', '•', '●')):
            raise HTTPException(422, '输入的是遮罩后的密钥，请点击平台的“复制”取得完整密钥，不要复制带星号或圆点的显示文字')
    parsed = urlsplit(body.base_url.strip())
    if parsed.scheme != 'https' or not parsed.hostname or parsed.username or parsed.password or parsed.query or parsed.fragment:
        raise HTTPException(422, 'API 地址必须是有效的 HTTPS 地址，且不能包含凭据、参数或片段')
    if body.reviewer_base_url.strip():
        reviewer = urlsplit(body.reviewer_base_url.strip())
        if reviewer.scheme != 'https' or not reviewer.hostname or reviewer.username or reviewer.password or reviewer.query or reviewer.fragment:
            raise HTTPException(422, '评审 API 地址必须为有效的 HTTPS 地址，且不能包含凭据或参数')
    current = settings()
    previous_host = urlsplit(current['base_url']).netloc.lower()
    if current.get('api_key') and previous_host != parsed.netloc.lower() and not (body.api_key and body.api_key.strip()):
        raise HTTPException(422, '切换 API 服务时请填写新平台的密钥，现有密钥不会发送到另一家服务商')
    old_review_host = urlsplit(current.get('reviewer_base_url') or current['base_url']).netloc.lower()
    new_review_host = urlsplit(body.reviewer_base_url.strip() or body.base_url.strip()).netloc.lower()
    if current.get('reviewer_api_key') and old_review_host != new_review_host and not (body.reviewer_api_key and body.reviewer_api_key.strip()):
        raise HTTPException(422, '切换评审服务时请填写新平台的评审密钥，原评审密钥不会发送到另一家服务商')
    update = body.model_dump(exclude={'api_key', 'reviewer_api_key'})
    update['base_url'] = body.base_url.strip().rstrip('/')
    update['reviewer_base_url'] = body.reviewer_base_url.strip().rstrip('/')
    current.update(update)
    if body.api_key and body.api_key.strip():
        current['api_key'] = body.api_key.strip()
    if body.reviewer_api_key and body.reviewer_api_key.strip():
        current['reviewer_api_key'] = body.reviewer_api_key.strip()
    if (current['reviewer_base_url'] and
            urlsplit(current['reviewer_base_url']).netloc != urlsplit(current['base_url']).netloc and
            not current['reviewer_api_key']):
        raise HTTPException(422, '独立评审服务需要自己的 API 密钥，不能将写作服务的密钥发送给另一家服务商')
    store().set_config('settings', current)
    return public_settings()


app.mount('/static', StaticFiles(directory=ROOT / 'static'), name='assets')
app.mount('/', StaticFiles(directory=ROOT / 'static', html=True), name='frontend')
